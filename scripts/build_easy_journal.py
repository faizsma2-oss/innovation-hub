"""Builds the easy journal: a spreadsheet with the journal.csv columns, calculations filled in, and a summary.

SIMULATION ONLY. Not for real money, and not a signal. Run from the repo root:
    .venv/bin/python scripts/build_easy_journal.py
It writes journal/xauusd-easy-journal.xlsx.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

CSV_COLUMNS = ["date", "session", "entry_time_local", "market", "setup", "signal_score", "direction", "entry",
               "stop", "target", "exit", "exit_reason", "result_R", "costs_usd", "simulated_size", "rule_followed",
               "actual_entry", "actual_exit", "notes", "source"]
HAND = ["rule_followed", "actual_entry", "actual_exit", "notes"]
AUTO = ["auto size (oz)", "result ($)", "cumulative R", "peak R", "drawdown R", "losing streak"]
LAST = 1000  # formulas and summary look this far down

HAND_FILL = PatternFill("solid", fgColor="FFF4CC")
AUTO_FILL = PatternFill("solid", fgColor="EFEFEF")
HEAD_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def build(path, rows=()):
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    journal = wb.create_sheet("Journal")
    summary = wb.create_sheet("Summary")
    settings = wb.create_sheet("Settings")

    # Settings: the limits the simulation uses. Confirmed by the user on 26 October 2026.
    settings.append(["Setting", "Value", "Note"])
    settings.append(["Simulated balance ($)", 10000, "Confirmed 26 Oct 2026. Keep it fixed during the test."])
    settings.append(["Risk per trade (fraction)", 0.005, "0.5%, confirmed. Keep it fixed during the test."])
    settings.append(["Risk amount per trade ($)", "=B2*B3", "Calculated. Used for every size and dollar figure."])
    settings.append(["Daily loss cap ($)", 100, "Kept at $100. Two full stops on one day cost about $103 with costs."])
    for cell in ("A1", "B1", "C1"):
        settings[cell].font = HEAD_FONT
    settings["B3"].number_format = "0.0%"
    settings.column_dimensions["A"].width = 30
    settings.column_dimensions["B"].width = 14
    settings.column_dimensions["C"].width = 70

    # Journal: the journal.csv columns (A to T), then the calculated columns (U to Z).
    headers = CSV_COLUMNS + AUTO
    journal.append(headers)
    for i, name in enumerate(headers, start=1):
        cell = journal.cell(row=1, column=i)
        cell.font = HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.fill = HAND_FILL if name in HAND else (AUTO_FILL if name in AUTO else PatternFill())
    for r_index, row in enumerate(rows, start=2):
        for c_index, value in enumerate(row, start=1):
            journal.cell(row=r_index, column=c_index, value=value)
    for r in range(2, LAST + 1):
        journal[f"U{r}"] = f'=IF(OR(H{r}="",I{r}=""),"",ROUND(Settings!$B$4/ABS(H{r}-I{r}),2))'
        journal[f"V{r}"] = f'=IF(M{r}="","",ROUND(M{r}*Settings!$B$4,2))'
        journal[f"W{r}"] = f'=IF(M{r}="","",SUM($M$2:M{r}))'
        journal[f"X{r}"] = f'=IF(M{r}="","",MAX(0,N(X{r - 1}),W{r}))'
        journal[f"Y{r}"] = f'=IF(M{r}="","",X{r}-W{r})'
        journal[f"Z{r}"] = f'=IF(M{r}="","",IF(M{r}<0,N(Z{r - 1})+1,0))'
        for col in ("P", "Q", "R", "S"):
            journal[f"{col}{r}"].fill = HAND_FILL
        for col in ("U", "V", "W", "X", "Y", "Z"):
            journal[f"{col}{r}"].fill = AUTO_FILL
        journal[f"S{r}"].alignment = WRAP
        journal[f"M{r}"].number_format = "+0.000;-0.000"
        journal[f"V{r}"].number_format = '"$"+0.00;"-$"0.00'
        for col in ("H", "I", "J", "K", "Q", "R"):
            journal[f"{col}{r}"].number_format = "0.00"
    journal.freeze_panes = "B2"
    journal.row_dimensions[1].height = 42
    widths = {"A": 12, "B": 11, "C": 11, "D": 9, "E": 40, "F": 10, "G": 9, "H": 11, "I": 11, "J": 11, "K": 11,
              "L": 11, "M": 11, "N": 10, "O": 12, "P": 13, "Q": 13, "R": 13, "S": 36, "T": 30,
              "U": 13, "V": 12, "W": 13, "X": 11, "Y": 13, "Z": 13}
    for col, width in widths.items():
        journal.column_dimensions[col].width = width

    validation = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    validation.error = "Type yes or no, or leave it blank if you didn't take the session."
    journal.add_data_validation(validation)
    validation.add(f"P2:P{LAST}")
    journal.conditional_formatting.add(f"M2:M{LAST}", CellIsRule(operator="greaterThan", formula=["0"],
                                                                 font=Font(color="2D6A3C")))
    journal.conditional_formatting.add(f"M2:M{LAST}", CellIsRule(operator="lessThan", formula=["0"],
                                                                 font=Font(color="8C2F2F")))

    # Summary: everything calculated from the journal, with a plain-English meaning beside each figure.
    # Each row is looked up by key, so every formula points at the right cell.
    keys = ["sessions", "wins", "winrate", "avg", "total", "total_usd", "avgwin", "avgloss", "dd_r", "dd_usd",
            "streak", "costs", "violations"]
    row = {key: i + 2 for i, key in enumerate(keys)}
    b = {key: f"B{row[key]}" for key in keys}
    summary.append(["Measure", "Value", "What it means"])
    lines = {
        "sessions": ("Sessions logged", f"=COUNT(Journal!M2:M{LAST})", "Each session with a result in the journal."),
        "wins": ("Wins", f'=COUNTIF(Journal!M2:M{LAST},">0")', "Sessions that finished with a positive R."),
        "winrate": ("Win rate", f'=IF({b["sessions"]}=0,"",{b["wins"]}/{b["sessions"]})',
                    "Wins as a share of sessions. Small samples swing a lot."),
        "avg": ("Average R per session", f'=IF({b["sessions"]}=0,"",AVERAGE(Journal!M2:M{LAST}))',
                "Average result in R, after costs."),
        "total": ("Total R", f"=SUM(Journal!M2:M{LAST})", "Sum of every logged result, in R."),
        "total_usd": ("Total $ (at the risk amount)", f"=ROUND({b['total']}*Settings!$B$4,2)",
                      "Total R times the risk amount per trade."),
        "avgwin": ("Average win (R)", f'=IFERROR(AVERAGEIF(Journal!M2:M{LAST},">0"),"none yet")',
                   "Average of the winning sessions."),
        "avgloss": ("Average loss (R)", f'=IFERROR(AVERAGEIF(Journal!M2:M{LAST},"<0"),"none yet")',
                    "Average of the losing sessions."),
        "dd_r": ("Max drawdown (R)", f"=MAX(Journal!Y2:Y{LAST})", "Largest fall from a peak in the running total."),
        "dd_usd": ("Max drawdown ($)", f"=ROUND({b['dd_r']}*Settings!$B$4,2)",
                   "The same fall, in dollars at the risk amount."),
        "streak": ("Longest losing streak", f"=MAX(Journal!Z2:Z{LAST})", "Most losing sessions in a row."),
        "costs": ("Spread and costs ($ per oz)", f"=SUM(Journal!N2:N{LAST})",
                  "Costs charged to the logged results, per ounce."),
        "violations": ("Rule violations (rule_followed = no)", f'=COUNTIF(Journal!P2:P{LAST},"no")',
                       "Sessions where a rule was broken. Log each one."),
    }
    for key in keys:
        label, formula, meaning = lines[key]
        summary.append([label, formula, meaning])
    for cell in ("A1", "B1", "C1"):
        summary[cell].font = HEAD_FONT
    formats = {"winrate": "0.0%", "avg": "+0.000;-0.000", "total": "+0.000;-0.000", "avgwin": "+0.000;-0.000",
               "avgloss": "+0.000;-0.000", "dd_r": "0.000", "total_usd": '"$"#,##0.00;"-$"#,##0.00',
               "dd_usd": '"$"#,##0.00', "costs": '"$"0.00'}
    for key, fmt in formats.items():
        summary[b[key]].number_format = fmt
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 14
    summary.column_dimensions["C"].width = 60

    # Read me: how to use the journal, in plain steps.
    text = [
        ("XAUUSD easy journal", True),
        ("SIMULATION ONLY. Not for real money. Not a signal. The rule has not yet been tested on data nobody has looked at.", True),
        ("", False),
        ("How to use it", True),
        ("1. On your Mac, run the journal command (see docs/xauusd-guide-complete.md). It writes data/journal.csv.", False),
        ("2. Open data/journal.csv in Excel or Numbers, so the columns split. Copy only the rows that are not in this sheet yet.", False),
        ("3. Paste them into the next empty row of the Journal sheet, starting at column A. Never paste over rows you've filled in.", False),
        ("4. Fill in the yellow columns by hand: rule_followed (yes or no), actual_entry, actual_exit and notes. Leave them blank if you didn't take that session.", False),
        ("5. The grey columns calculate themselves. Don't type in them.", False),
        ("6. The Summary sheet updates by itself.", False),
        ("", False),
        ("The rule the journal tracks", True),
        ("Session momentum on XAUUSD. London and New York sessions, 08:00 to 17:00 local time. Entry at 08:30 local time if the close is at least 0.6 × daily ATR above or below the pre-session start price and on the matching side of the 300-bar EMA. Stop 0.6 × daily ATR from entry. Target 2 × the stop distance. Exit at the stop, the target, or the 17:00 close, whichever comes first. At most one entry per session.", False),
        ("", False),
        ("The limits", True),
        ("Simulated balance $10,000. Risk $50 per trade (0.5%). Size = $50 ÷ the distance between entry and stop. Daily loss cap $100. Limits are on the Settings sheet and are fixed for the test.", False),
    ]
    for line, bold in text:
        readme.append([line])
        if bold:
            readme[f"A{readme.max_row}"].font = Font(bold=True, size=13 if line.startswith("XAUUSD") else 11)
        readme[f"A{readme.max_row}"].alignment = WRAP
    readme.column_dimensions["A"].width = 110

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = build(Path(__file__).resolve().parent.parent / "journal" / "xauusd-easy-journal.xlsx")
    print(f"Wrote {out}")
    sys.exit(0)
